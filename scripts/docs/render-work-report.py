#!/usr/bin/env python3
"""Render internal report Markdown as matching A4 Word/PDF using local Pango/Cairo."""
import argparse
import html
from pathlib import Path

import cairo
import markdown
from bs4 import BeautifulSoup
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
import report_pdf as base


def widths(table):
    cells = table.find('tr').find_all(['th', 'td'])
    first = cells[0].get_text().strip()
    count = len(cells)
    if count == 2:
        return [.27, .73]
    if count == 3:
        if first in {'Date', 'Period'}:
            return [.17, .62, .21]
        if first in {'Task', 'Activity', 'Description'}:
            return [.62, .21, .17]
        return [.28, .48, .24]
    if count == 4:
        return [.13, .36, .29, .22]
    return [1 / count] * count


base.proportions = widths


class Pdf(base.PdfDocument):
    def __init__(self, path, title, period, author='Thomas'):
        super().__init__(path)
        self.period = period
        self.surface.set_metadata(cairo.PDF_METADATA_TITLE, title)
        self.surface.set_metadata(cairo.PDF_METADATA_AUTHOR, author)

    def finish_page(self):
        self.color(base.BORDER)
        self.context.set_line_width(.5)
        self.context.move_to(self.margin, 803)
        self.context.line_to(self.width - self.margin, 803)
        self.context.stroke()
        label, _ = self.layout(html.escape(f'Toolkit ICT | {self.period} | Internal | Page {self.page}'), 510, 7)
        self.draw(label, self.margin, 812, '68736d')
        self.context.show_page()


def render(source, period, logo=None, author='Thomas'):
    logo = logo or ROOT / 'docs/assets/toolkit-logo.png'
    soup = BeautifulSoup(markdown.markdown(source.read_text(), extensions=['tables']), 'html.parser')
    nodes = [node for node in soup.children if getattr(node, 'name', None)]
    title = soup.find('h1').get_text() if soup.find('h1') else source.stem
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(17)
    section.left_margin = section.right_margin = Mm(15)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(9.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    for style, size in [('Heading 1', 22), ('Heading 2', 15), ('Heading 3', 11.5)]:
        document.styles[style].font.name = 'Arial'
        document.styles[style].font.size = Pt(size)
        document.styles[style].font.color.rgb = RGBColor.from_string(base.TEAL)
        document.styles[style].paragraph_format.keep_with_next = True
    document.add_picture(str(logo), width=Mm(42))
    pdf = Pdf(source.with_suffix('.pdf'), title, period, author)
    pdf.logo(logo)
    for node in nodes:
        if node.name in {'h1', 'h2', 'h3', 'p'}:
            paragraph = document.add_paragraph(style={'h1': 'Heading 1', 'h2': 'Heading 2', 'h3': 'Heading 3'}.get(node.name, 'Normal'))
            base.docx_inline(paragraph, node)
            pdf.paragraph(node)
        elif node.name in {'ul', 'ol'}:
            for item in node.find_all('li', recursive=False):
                paragraph = document.add_paragraph(style='List Bullet' if node.name == 'ul' else 'List Number')
                base.docx_inline(paragraph, item)
                block = BeautifulSoup('<p>• ' + item.decode_contents() + '</p>', 'html.parser').p
                pdf.paragraph(block)
        elif node.name == 'table':
            ws = widths(node)
            table = document.add_table(rows=0, cols=len(ws))
            table.style, table.autofit = 'Table Grid', False
            for column, fraction in zip(table.columns, ws):
                column.width = Mm(180 * fraction)
            for index, row in enumerate(node.find_all('tr')):
                cells = table.add_row().cells
                for cell, content, fraction in zip(cells, row.find_all(['th', 'td']), ws):
                    cell.width = Mm(180 * fraction)
                    paragraph = cell.paragraphs[0]
                    base.docx_inline(paragraph, content, bold=index == 0)
                    for run in paragraph.runs:
                        run.font.size = Pt(8)
                    if index == 0:
                        base.shade(cell, base.LIGHT)
                props = table.rows[-1]._tr.get_or_add_trPr()
                props.append(OxmlElement('w:cantSplit'))
                if index == 0:
                    props.append(OxmlElement('w:tblHeader'))
            document.add_paragraph()
            pdf.table(node)
    footer = section.footer.paragraphs[0]
    footer.add_run(f'Toolkit ICT | {period} | Internal | Page ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(7)
    document.core_properties.title = title
    document.core_properties.author = author
    document.save(source.with_suffix('.docx'))
    print(source.name, pdf.finish(), 'PDF pages; Word and PDF prepared')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('--period', default='22 September–23 October 2026')
    parser.add_argument('--logo', type=Path, help='Existing approved institution logo')
    parser.add_argument('--author', default='Thomas')
    args = parser.parse_args()
    render(args.source, args.period, args.logo, args.author)
