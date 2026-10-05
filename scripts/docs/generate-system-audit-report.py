#!/usr/bin/env python3
"""Render an illustrated system-audit Markdown report as Word/PDF using local Python and Pango/Cairo.

Uses the existing report's A4 geometry, Arial-compatible text, teal headings,
orange rule and light table headers. This route needs no LibreOffice runtime.
"""

import argparse
import html
from io import BytesIO
from pathlib import Path

import cairo
import gi
import markdown
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image

gi.require_version("Pango", "1.0")
gi.require_version("PangoCairo", "1.0")
from gi.repository import Pango, PangoCairo

REPO = Path(__file__).resolve().parents[2]
TEAL, INK, ORANGE, LIGHT, BORDER = "174a43", "26332f", "e46218", "dfe7e3", "c6ceca"


def markup(node):
    if isinstance(node, NavigableString):
        return html.escape(str(node))
    if node.name == "br":
        return "\n"
    inner = "".join(markup(child) for child in node.children)
    if node.name == "a" and node.get("href", "").startswith("https://"):
        return inner + " (" + html.escape(node["href"]) + ")"
    if node.name in {"strong", "b"}:
        return f"<b>{inner}</b>"
    if node.name in {"em", "i"}:
        return f"<i>{inner}</i>"
    return inner


def plain(node):
    return BeautifulSoup(markup(node).replace("\n", "<br/>"), "html.parser").get_text(separator="", strip=False)


def proportions(table):
    heads = [plain(cell) for cell in table.find("tr").find_all(["th", "td"])]
    if heads[0] == "Group":
        return [0.10, 0.13, 0.77]
    if heads[0] == "Ref":
        return [0.08, 0.92]
    if heads[0] == "Measure":
        return [0.36, 0.19, 0.45]
    if len(heads) == 2:
        return [0.27, 0.73]
    if len(heads) == 4:
        return [0.29, 0.20, 0.27, 0.24]
    return [0.27, 0.43, 0.30]


def docx_inline(paragraph, node, bold=False, italic=False):
    if isinstance(node, NavigableString):
        run = paragraph.add_run(str(node))
        run.bold, run.italic = bold, italic
    elif node.name == "br":
        paragraph.add_run().add_break()
    elif node.name == "a" and node.get("href", "").startswith("https://"):
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        relationship = paragraph.part.relate_to(node["href"], RT.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), relationship)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        colour = OxmlElement("w:color"); colour.set(qn("w:val"), TEAL)
        underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single")
        properties.extend([colour, underline]); run.append(properties)
        text = OxmlElement("w:t"); text.text = node.get_text()
        run.append(text); link.append(run); paragraph._p.append(link)
    else:
        for child in node.children:
            docx_inline(paragraph, child, bold or node.name in {"strong", "b"}, italic or node.name in {"em", "i"})


def shade(cell, fill):
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(item)


def word_document(nodes, path, source_dir):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(17)
    section.left_margin = section.right_margin = Mm(15)
    base = document.styles["Normal"]
    base.font.name, base.font.size = "Arial", Pt(9.5)
    base.font.color.rgb = RGBColor.from_string(INK)
    base.paragraph_format.space_after, base.paragraph_format.line_spacing = Pt(7), 1.15
    for name, size, color in [("Heading 1", 23, TEAL), ("Heading 2", 15, TEAL), ("Heading 3", 11.5, "4e592e")]:
        style = document.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
    document.add_picture(str((REPO / "docs/assets/toolkit-logo.png" if (REPO / "docs/assets/toolkit-logo.png").exists() else REPO / "wp-content/themes/eduma-child/assets/images/toolkit-logo.png")), width=Mm(46.5))
    for node in nodes:
        if node.find('img'):
            image = node.find('img')
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(source_dir / image['src']), width=Mm(180))
            caption = document.add_paragraph(image.get('alt', ''), style='Caption')
            caption.paragraph_format.space_after = Pt(12)
            continue
        if node.name in {"h1", "h2", "h3", "p"}:
            paragraph = document.add_paragraph(style={"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}.get(node.name, "Normal"))
            docx_inline(paragraph, node)
            if node.name == "h1":
                borders = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                for key, value in {"val": "single", "sz": "18", "color": ORANGE, "space": "7"}.items():
                    bottom.set(qn("w:" + key), value)
                borders.append(bottom)
                paragraph._p.get_or_add_pPr().append(borders)
        elif node.name == "table":
            rows = node.find_all("tr")
            table = document.add_table(rows=0, cols=len(proportions(node)))
            table.style, table.autofit = "Table Grid", False
            widths = [Mm(180 * fraction) for fraction in proportions(node)]
            for column, width in zip(table.columns, widths):
                column.width = width
            for index, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, source, width in zip(cells, row.find_all(["th", "td"]), widths):
                    cell.width = width
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(4)
                    docx_inline(paragraph, source, bold=index == 0)
                    for run in paragraph.runs:
                        run.font.size = Pt(7.7)
                    if index == 0:
                        shade(cell, LIGHT)
                props = table.rows[-1]._tr.get_or_add_trPr()
                props.append(OxmlElement("w:cantSplit"))
                if index == 0:
                    props.append(OxmlElement("w:tblHeader"))
            document.add_paragraph().paragraph_format.space_after = Pt(2)
    footer = section.footer.paragraphs[0]
    footer.add_run("Toolkit ICT | System audit | 5 October 2026 | Internal | Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(7)
    document.core_properties.title = "Toolkit Africa System Audit Remediation Report — 5 October 2026"
    document.core_properties.subject = "Verified findings, deployed corrections and evidence"
    document.core_properties.author = "The Toolkit for Skills and Innovation"
    document.save(path)


class PdfDocument:
    width, height, margin, bottom = 595.276, 841.89, 42.52, 788

    def __init__(self, path):
        self.surface = cairo.PDFSurface(str(path), self.width, self.height)
        self.surface.set_metadata(cairo.PDF_METADATA_TITLE, "Toolkit Africa System Audit Remediation Report — 5 October 2026")
        self.context = cairo.Context(self.surface)
        self.y, self.page = 48, 1

    def color(self, value):
        self.context.set_source_rgb(*(int(value[i:i + 2], 16) / 255 for i in (0, 2, 4)))

    def layout(self, text, width, size, bold=False):
        layout = PangoCairo.create_layout(self.context)
        PangoCairo.context_set_resolution(layout.get_context(), 72)
        layout.set_font_description(Pango.FontDescription(f"Liberation Sans {'Bold ' if bold else ''}{size}"))
        layout.set_width(round(width * Pango.SCALE))
        layout.set_wrap(Pango.WrapMode.WORD_CHAR)
        layout.set_spacing(round(size * 0.18 * Pango.SCALE))
        layout.set_markup(text, -1)
        return layout, layout.get_size()[1] / Pango.SCALE

    def draw(self, layout, x, y, color=INK):
        self.color(color)
        self.context.move_to(x, y)
        PangoCairo.show_layout(self.context, layout)

    def finish_page(self):
        self.color(BORDER)
        self.context.set_line_width(0.5)
        self.context.move_to(self.margin, 803)
        self.context.line_to(self.width - self.margin, 803)
        self.context.stroke()
        label, _ = self.layout(f"Toolkit ICT | System audit | 5 October 2026 | Internal | Page {self.page}", 510, 7)
        self.draw(label, self.margin, 812, "68736d")
        self.context.show_page()

    def new_page(self):
        self.finish_page()
        self.page += 1
        self.y = 48

    def logo(self):
        with Image.open((REPO / "docs/assets/toolkit-logo.png" if (REPO / "docs/assets/toolkit-logo.png").exists() else REPO / "wp-content/themes/eduma-child/assets/images/toolkit-logo.png")) as source:
            buffer = BytesIO()
            source.convert("RGBA").save(buffer, "PNG")
        buffer.seek(0)
        logo = cairo.ImageSurface.create_from_png(buffer)
        ratio = 132 / logo.get_width()
        self.context.save()
        self.context.translate(self.margin, self.y)
        self.context.scale(ratio, ratio)
        self.context.set_source_surface(logo, 0, 0)
        self.context.paint()
        self.context.restore()
        self.y += logo.get_height() * ratio + 10

    def paragraph(self, node):
        heading = node.name.startswith("h")
        size = {"h1": 23, "h2": 15, "h3": 11.5}.get(node.name, 9.5)
        before = 12 if heading else 0
        layout, height = self.layout(markup(node), self.width - self.margin * 2, size, heading)
        if self.y + before + height + (35 if heading else 0) > self.bottom:
            self.new_page()
        self.y += before
        self.draw(layout, self.margin, self.y, TEAL if heading else INK)
        self.y += height + 7
        if node.name == "h1":
            self.color(ORANGE)
            self.context.rectangle(self.margin, self.y, self.width - self.margin * 2, 3)
            self.context.fill()
            self.y += 12

    def table(self, node):
        widths = [(self.width - self.margin * 2) * fraction for fraction in proportions(node)]
        source_rows = node.find_all("tr")

        def prepare(row, header):
            layouts = [self.layout(markup(cell), width - 10, 7.7, header)
                       for cell, width in zip(row.find_all(["th", "td"]), widths)]
            return layouts, max(height for _, height in layouts) + 10

        def draw_row(layouts, height, header=False):
            x = self.margin
            for (layout, _), width in zip(layouts, widths):
                self.color(LIGHT if header else "ffffff")
                self.context.rectangle(x, self.y, width, height)
                self.context.fill_preserve()
                self.color(BORDER)
                self.context.set_line_width(0.5)
                self.context.stroke()
                self.draw(layout, x + 5, self.y + 5)
                x += width
            self.y += height

        header_layouts, header_height = prepare(source_rows[0], True)
        first_height = prepare(source_rows[1], False)[1] if len(source_rows) > 1 else 0
        if self.y + header_height + first_height > self.bottom:
            self.new_page()
        draw_row(header_layouts, header_height, True)
        for row in source_rows[1:]:
            layouts, height = prepare(row, False)
            if self.y + height > self.bottom:
                self.new_page()
                draw_row(header_layouts, header_height, True)
            draw_row(layouts, height)
        self.y += 11

    def evidence_image(self, node, source_dir):
        image = node.find('img')
        surface = cairo.ImageSurface.create_from_png(str(source_dir / image['src']))
        width = self.width - self.margin * 2
        ratio = min(width / surface.get_width(), 315 / surface.get_height())
        height = surface.get_height() * ratio
        caption, caption_height = self.layout(html.escape(image.get('alt', '')), width, 8)
        if self.y + height + caption_height + 24 > self.bottom:
            self.new_page()
        self.context.save()
        self.context.translate(self.margin + (width - surface.get_width() * ratio) / 2, self.y)
        self.context.scale(ratio, ratio)
        self.context.set_source_surface(surface, 0, 0)
        self.context.paint()
        self.context.restore()
        self.y += height + 6
        self.draw(caption, self.margin, self.y, '58685d')
        self.y += caption_height + 14

    def finish(self):
        self.finish_page()
        self.surface.finish()
        return self.page


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    args = parser.parse_args()
    body = BeautifulSoup(markdown.markdown(args.source.read_text(), extensions=["tables"]), "html.parser")
    nodes = [node for node in body.children if getattr(node, "name", None)]
    word_document(nodes, args.source.with_suffix(".docx"), args.source.parent)
    pdf = PdfDocument(args.source.with_suffix(".pdf"))
    pdf.logo()
    for node in nodes:
        if node.find("img"):
            pdf.evidence_image(node, args.source.parent)
        elif node.name == "table":
            pdf.table(node)
        else:
            pdf.paragraph(node)
    print(f"Wrote Word and {pdf.finish()}-page PDF: {args.source.stem}")


if __name__ == "__main__":
    main()
